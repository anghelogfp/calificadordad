"""Generadores de salida del reporte de calificación (CSV / XLSX / PDF)."""
from __future__ import annotations

import csv
import datetime as _dt
from itertools import groupby

CSV_HEADER = [
    "N°", "Salón", "DNI", "Apellidos y Nombres", "Escuela Profesional",
    "Facultad", "Plaza", "Correctas", "Incorrectas", "En blanco", "Anuladas",
    "Aciertos", "Puntaje /30", "Puntaje /20", "Condición",
]

_VACIA = ("", "\\N", "\\n", "NULL", None)


def _plaza(row: dict) -> str:
    plaza = row.get("plaza", "") or ""
    extra = row.get("plaza_adicional", "") or ""
    if extra not in _VACIA:
        return f"{plaza} / {extra}"
    return plaza


def _fila(row: dict) -> list:
    return [
        row.get("n", ""), row.get("salon", ""), row["dni"], row["nombre"],
        row["escuela"], row.get("facultad", ""), _plaza(row),
        row["correctas"], row["incorrectas"], row["blancos"], row["anuladas"],
        row["aciertos"], row["puntaje30"], f'{row["puntaje20"]:.3f}',
        row["condicion"],
    ]


def write_csv(path, report) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(CSV_HEADER)
        for row in report.resultados:
            w.writerow(_fila(row))


def write_xlsx(path, report) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    verde = PatternFill("solid", fgColor="C6EFCE")
    rojo = PatternFill("solid", fgColor="FFC7CE")
    negrita = Font(bold=True)

    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Resultados"
    ws.append(CSV_HEADER)
    for c in ws[1]:
        c.font = negrita
    cond_col = CSV_HEADER.index("Condición") + 1
    for row in report.resultados:
        ws.append(_fila(row))
        ws.cell(row=ws.max_row, column=cond_col).fill = (
            verde if row["condicion"] == "APTO" else rojo)
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("Resumen")
    ws.append(["Escuela Profesional", "Postulantes", "APTO", "NO APTO"])
    for c in ws[1]:
        c.font = negrita
    for e in report.resumen:
        ws.append([e["escuela"], e["postulantes"], e["aptos"], e["no_aptos"]])
    ws.append([
        "TOTAL",
        sum(e["postulantes"] for e in report.resumen),
        sum(e["aptos"] for e in report.resumen),
        sum(e["no_aptos"] for e in report.resumen),
    ])
    for c in ws[ws.max_row]:
        c.font = negrita

    ws = wb.create_sheet("Incidencias")
    ws.append(["Tipo", "DNI", "Litho", "Nombre", "Escuela", "Detalle"])
    for c in ws[1]:
        c.font = negrita
    for i in report.incidencias:
        ws.append([i.get("tipo", ""), i.get("dni", ""), i.get("litho", ""),
                   i.get("nombre", ""), i.get("escuela", ""), i.get("detalle", "")])

    for sheet in wb.worksheets:
        for col in sheet.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None),
                        default=10)
            sheet.column_dimensions[col[0].column_letter].width = min(width + 2, 60)

    wb.save(path)


def write_pdf(path, report, *, titulo: str, subtitulo: str,
              entidad: str = "UNIVERSIDAD NACIONAL DEL ALTIPLANO DE PUNO",
              comision: str = "COMISIÓN DE CONCURSO PÚBLICO DE CÁTEDRA 2026-II",
              umbral: int | None = None, logo=None) -> None:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (BaseDocTemplate, Frame, Image, PageBreak,
                                    Paragraph, PageTemplate, Spacer, Table,
                                    TableStyle)

    page = A4  # retrato
    styles = getSampleStyleSheet()
    h_ent = ParagraphStyle("ent", parent=styles["Title"], fontSize=12, spaceAfter=2,
                           alignment=1, leading=14)
    h_com = ParagraphStyle("com", parent=styles["Normal"], fontSize=9, alignment=1,
                           leading=11)
    h_tit = ParagraphStyle("tit", parent=styles["Normal"], fontSize=10, alignment=1,
                           spaceBefore=8, spaceAfter=2, fontName="Helvetica-Bold")
    h_esc = ParagraphStyle("esc", parent=styles["Normal"], fontSize=15, alignment=1,
                           leading=18, fontName="Helvetica-Bold",
                           textColor=colors.HexColor("#1A1A1A"))

    generado = _dt.datetime.now().strftime("%Y/%m/%d %H:%M:%S")
    nota = f"APTO = {umbral} aciertos o más (nota ≥ 11)." if umbral else ""

    def _footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(15 * mm, 10 * mm, generado)
        canvas.drawRightString(page[0] - 15 * mm, 10 * mm, f"Página {doc.page}")
        if nota:
            canvas.drawCentredString(page[0] / 2, 10 * mm, nota)
        canvas.restoreState()

    frame = Frame(15 * mm, 15 * mm, page[0] - 30 * mm, page[1] - 30 * mm, id="f")
    doc = BaseDocTemplate(str(path), pagesize=page,
                          pageTemplates=[PageTemplate(id="p", frames=[frame],
                                                      onPage=_footer)])

    import os
    logo_path = str(logo) if logo is not None and os.path.exists(str(logo)) else None

    # ancho útil retrato A4 = 180 mm
    col_w = [10 * mm, 24 * mm, 64 * mm, 20 * mm, 18 * mm, 20 * mm, 24 * mm]
    head = ["N°", "DNI", "APELLIDOS Y NOMBRES", "ACIERTOS", "PTS/30", "NOTA/20",
            "CONDICIÓN"]

    _ = titulo  # se mantiene el parámetro por compatibilidad; ya no se imprime

    def _masthead():
        textos = [Paragraph(entidad, h_ent), Paragraph(comision, h_com),
                  Paragraph(subtitulo, h_tit)]
        if logo_path is None:
            return textos

        def _logo():
            return Image(logo_path, width=18 * mm, height=18 * mm,
                         kind="proportional")

        # 3 columnas: insignia | texto centrado en la página | insignia
        mh = Table([[_logo(), textos, _logo()]],
                   colWidths=[22 * mm, 136 * mm, 22 * mm])
        mh.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("ALIGN", (0, 0), (0, 0), "LEFT"),
            ("ALIGN", (1, 0), (1, 0), "CENTER"),
            ("ALIGN", (2, 0), (2, 0), "RIGHT"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ("TOPPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return [mh]

    def _escuela_band(escuela):
        band = Table([[Paragraph(escuela.upper(), h_esc)]], colWidths=[180 * mm])
        band.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#E9EEF5")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#9BB2CC")),
            ("TOPPADDING", (0, 0), (-1, -1), 7),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ]))
        return band

    story = []
    grupos = groupby(report.resultados, key=lambda r: (r["escuela_id"], r["escuela"]))
    first = True
    for (_eid, escuela), filas in grupos:
        if not first:
            story.append(PageBreak())
        first = False
        story.extend(_masthead())
        story.append(Spacer(0, 4 * mm))
        story.append(_escuela_band(escuela))
        story.append(Spacer(0, 4 * mm))

        data = [head]
        estilo = [
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E6E6E6")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F4F4F4")]),
            ("LINEBELOW", (0, 0), (-1, 0), 0.75, colors.black),
            ("LINEBELOW", (0, -1), (-1, -1), 0.5, colors.black),
            ("ALIGN", (3, 0), (-1, -1), "CENTER"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i, r in enumerate(filas, start=1):
            data.append([str(i), r["dni"], r["nombre"], str(r["aciertos"]),
                         str(r["puntaje30"]), f'{r["puntaje20"]:.2f}',
                         r["condicion"]])
            if r["condicion"] == "NO APTO":
                estilo.append(("TEXTCOLOR", (-1, i), (-1, i),
                               colors.HexColor("#B00000")))
        t = Table(data, colWidths=col_w, repeatRows=1)
        t.setStyle(TableStyle(estilo))
        story.append(t)

    if not story:
        story.append(Paragraph("Sin resultados para mostrar.", styles["Normal"]))
    doc.build(story)
